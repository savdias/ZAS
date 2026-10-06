from datetime import date
from io import BytesIO
import importlib.util
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
import pytest

import dados


def arquivo_excel(linhas, nome_aba="Barragens"):
    workbook = Workbook()
    ws = workbook.active
    ws.title = nome_aba
    for linha in linhas:
        ws.append(linha)
    return workbook


def buffer_excel(workbook):
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)
    return arquivo


@pytest.mark.parametrize("valor, esperado", [
    (1234, 1234.0), ("1.234", 1234.0), ("1.234,56", 1234.56),
    ("1,234.56", 1234.56), ("1234,5", 1234.5), ("12.5", 12.5),
    ("0.125", 0.125), (" 1\u00a0234,5 ", 1234.5), ("(1.234,5)", -1234.5),
    (None, None), ("", None), ("Não", None), ("NaN", None), (float("inf"), None),
    (True, None), ("1,2,3", None), ("1.2.3", None),
])
def test_numeros_brasileiros_e_valores_invalidos(valor, esperado):
    assert dados.converter_numero(valor) == esperado


def test_formatacao_nao_trunca_decimais():
    assert dados.formatar_numero(1234) == "1.234"
    assert dados.formatar_numero(1234.5) == "1.234,5"
    assert dados.formatar_numero(float("nan")) == "—"


def test_seleciona_aba_do_dominio_em_vez_da_mais_larga():
    workbook = arquivo_excel([
        ["Relatório anual"],
        ["NMBarragem", "UF", "NMMunicipio", "TrabalhadoresZAS"],
        ["B1", "MG", "M1", "1.234"],
    ])
    resumo = workbook.create_sheet("Resumo")
    resumo.append(["a", "b", "c", "d", "e", "f"])
    resumo.append([1, 2, 3, 4, 5, 6])
    cabecalhos, registros = dados.carregar_planilha(buffer_excel(workbook))
    assert cabecalhos == ["NMBarragem", "UF", "NMMunicipio", "TrabalhadoresZAS"]
    assert registros[0]["NMBarragem"] == "B1"
    assert dados.calcular_indicadores(registros)["trabalhadores"] == "1.234"


def test_linha_de_dados_com_palavras_do_dominio_nao_vira_cabecalho():
    workbook = arquivo_excel([
        ["Nome da Barragem", "UF", "Município"],
        ["Barragem Município Estado ZAS comunidade empreendedor mineradora", "MG", "Nova Lima"],
    ])
    assert dados.descobrir_cabecalho(workbook.active) == 1
    _, registros = dados.carregar_planilha(buffer_excel(workbook))
    assert len(registros) == 1


def test_preserva_colunas_duplicadas_e_datas_serializaveis():
    workbook = arquivo_excel([
        ["Barragem", "UF", "UF", "Data"],
        ["B1", "MG", "SP", date(2026, 1, 2)],
        ["   ", None, None, None],
    ])
    cabecalhos, registros = dados.carregar_planilha(buffer_excel(workbook))
    assert cabecalhos == ["Barragem", "UF", "UF (2)", "Data"]
    assert len(registros) == 1
    assert registros[0]["UF (2)"] == "SP"
    assert registros[0]["Data"].startswith("2026-01-02")
    json.dumps(registros, allow_nan=False)


def test_rejeita_arquivo_sem_cabecalhos_reconhecidos_e_fecha_workbook(monkeypatch):
    workbook = arquivo_excel([["A", "B"], [1, 2]])
    fechamentos = []
    monkeypatch.setattr(workbook, "close", lambda: fechamentos.append(True))
    monkeypatch.setattr(dados, "load_workbook", lambda *a, **kw: workbook)
    with pytest.raises(ValueError, match="cabeçalhos reconhecidos"):
        dados.carregar_planilha(BytesIO())
    assert fechamentos == [True]


def test_conta_ids_e_municipios_homonimos_em_estados_distintos():
    registros = [
        {"IDBarragem": 1, "Barragem": "B1", "UF": "MG", "Município": "Santa Rita"},
        {"IDBarragem": 2, "Barragem": "B1", "UF": "SP", "Município": "Santa Rita"},
        {"IDBarragem": 2, "Barragem": "B1", "UF": "SP", "Município": "Santa Rita"},
    ]
    indicadores = dados.calcular_indicadores(registros)
    assert indicadores["barragens"] == "2"
    assert indicadores["municipios"] == "2"
    assert indicadores["estados"] == "2"
    assert indicadores["registros"] == "3"


def test_nome_barragem_com_localizacao_e_colunas_sem_reuso():
    registros = [
        {"Barragem": "B1", "UF": "MG", "Município": "M1"},
        {"Barragem": "B1", "UF": "MG", "Município": "M2"},
    ]
    assert dados.calcular_indicadores(registros)["barragens"] == "2"
    colunas = dados.identificar_colunas(["IDBarragem", "Município da Barragem"])
    assert colunas["id"] == "IDBarragem"
    assert colunas["municipios"] == "Município da Barragem"
    assert colunas["barragens"] is None


def test_trabalhadores_sim_nao_sao_contagem_de_pessoas():
    registros = [
        {"IDBarragem": 1, "NMBarragem": "Barragem 01", "TrabalhadoresZAS": "Sim"},
        {"IDBarragem": 2, "NMBarragem": "Barragem 01", "TrabalhadoresZAS": "Não"},
        {"IDBarragem": 3, "NMBarragem": "Barragem 02", "TrabalhadoresZAS": "Não informado"},
    ]
    assert dados.calcular_indicadores(registros)["barragens"] == "3"
    assert dados.calcular_indicadores(registros)["trabalhadores"] == "—"
    assert dados.contar_respostas_sim(registros, "TrabalhadoresZAS") == 1
    assert dados.contar_respostas_sim([{"x": "Não"}], "x") == 0
    assert dados.contar_respostas_sim([{"x": "Não informado"}], "x") is None


def test_dados_ausentes_nao_sao_zero():
    registros = [{"Barragem": "B1", "TrabalhadoresZAS": None, "UF": "   "}]
    indicadores = dados.calcular_indicadores(registros)
    assert indicadores["trabalhadores"] == "—"
    assert indicadores["municipios"] == "—"
    assert dados.valores_unicos("UF", registros) == set()
    assert dados.somar_coluna("TrabalhadoresZAS", [{"TrabalhadoresZAS": 0}]) == 0


def test_localiza_planilha_ignore_temporario_e_aceita_extensao_maiuscula(tmp_path, monkeypatch):
    (tmp_path / "~$Barragem.xlsx").touch()
    (tmp_path / "outro.xlsx").touch()
    escolhida = tmp_path / "Lista ANM.XLSX"
    escolhida.touch()
    monkeypatch.setattr(dados, "PASTA_DADOS", tmp_path)
    assert dados.localizar_planilha() == escolhida


def test_planilha_de_coordenadas_nao_substitui_lista_principal(tmp_path, monkeypatch):
    (tmp_path / "coordenadas_barragens_preenchidas.xlsx").touch()
    principal = tmp_path / "Lista Barragens AIR ANM 2026.xlsx"
    principal.touch()
    monkeypatch.setattr(dados, "PASTA_DADOS", tmp_path)
    assert dados.localizar_planilha() == principal
    principal.unlink()
    with pytest.raises(FileNotFoundError):
        dados.localizar_planilha()


def test_upload_nao_altera_dados_globais():
    registros_originais = dados.REGISTROS
    cabecalhos_originais = dados.CABECALHOS
    workbook = arquivo_excel([["Barragem"], ["Importada"]])
    _, registros = dados.carregar_planilha(buffer_excel(workbook))
    assert registros == [{"Barragem": "Importada"}]
    assert dados.REGISTROS is registros_originais
    assert dados.CABECALHOS is cabecalhos_originais


def test_import_sem_planilha_continua_com_indicadores_vazios(tmp_path):
    modulo = tmp_path / "dados.py"
    modulo.write_text(Path(dados.__file__).read_text(), encoding="utf-8")
    spec = importlib.util.spec_from_file_location("dados_sem_planilha", modulo)
    sem_planilha = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(sem_planilha)
    assert sem_planilha.REGISTROS == []
    assert sem_planilha.CABECALHOS == []
    assert sem_planilha.INDICADORES["registros"] == "0"
    assert sem_planilha.INDICADORES["trabalhadores"] == "—"
    assert "Nenhuma planilha XLSX" in sem_planilha.ERRO_DADOS


def test_respostas_sim_deduplicam_id_e_preservam_homonimos():
    registros = [
        {"IDBarragem": 1, "Barragem": "B1", "TrabalhadoresZAS": "Sim"},
        {"IDBarragem": 1, "Barragem": "B1", "TrabalhadoresZAS": "Sim"},
        {"IDBarragem": 2, "Barragem": "B1", "TrabalhadoresZAS": "Sim"},
    ]
    assert dados.contar_respostas_sim(registros, "TrabalhadoresZAS") == 2
    sem_id = [
        {"Barragem": "B1", "UF": "MG", "Município": "M1", "Comunidade ZAS": "Sim"},
        {"Barragem": "B1", "UF": "MG", "Município": "M1", "Comunidade ZAS": "Sim"},
        {"Barragem": "B1", "UF": "SP", "Município": "M1", "Comunidade ZAS": "Sim"},
    ]
    assert dados.contar_respostas_sim(sem_id, "Comunidade ZAS") == 2


def test_quantidades_numericas_nao_sao_respostas_sim_nao():
    for valor in (0, 1, 1.0, "0", "1", "1,0"):
        registros = [{"Barragem": "B1", "TrabalhadoresZAS": valor}]
        assert dados.contar_respostas_sim(registros, "TrabalhadoresZAS") is None
        assert dados.calcular_indicadores(registros)["trabalhadores"] != "—"


def test_rejeita_planilha_sem_id_ou_nome_de_barragem():
    workbook = arquivo_excel([["NMMineradora"], ["Empresa A"]])
    with pytest.raises(ValueError, match="ID ou o nome da barragem"):
        dados.carregar_planilha(buffer_excel(workbook))


def test_erro_xml_local_nao_impede_carregar_upload(tmp_path):
    modulo = tmp_path / "dados.py"
    modulo.write_text(Path(dados.__file__).read_text(), encoding="utf-8")
    pasta = tmp_path / "dados"
    pasta.mkdir()
    valido = buffer_excel(arquivo_excel([["Barragem"], ["B1"]]))
    with ZipFile(valido) as original, ZipFile(pasta / "corrompido.xlsx", "w") as destino:
        for info in original.infolist():
            conteudo = original.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                conteudo = conteudo[:-12]
            destino.writestr(info, conteudo)
    spec = importlib.util.spec_from_file_location("dados_xml_invalido", modulo)
    sem_planilha = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(sem_planilha)
    assert sem_planilha.REGISTROS == []
    assert sem_planilha.ERRO_DADOS
    valido.seek(0)
    _, registros = sem_planilha.carregar_planilha(valido)
    assert registros == [{"Barragem": "B1"}]
