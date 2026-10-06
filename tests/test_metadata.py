from openpyxl import Workbook, load_workbook

import coordenadas
from coordenadas import carregar_coordenadas, carregar_metadados, enriquecer_registros


def salvar_referencia(tmp_path, linhas, cabecalho=None):
    livro = Workbook()
    aba = livro.active
    aba.title = "Coordenadas"
    aba.append(cabecalho or ["IDBarragem", "Status", "Fonte", "Observação"])
    for linha in linhas:
        aba.append(linha)
    caminho = tmp_path / "referencia.xlsx"
    livro.save(caminho)
    livro.close()
    return caminho


def test_metadados_ausentes_nao_bloqueiam_leitura(tmp_path, monkeypatch):
    monkeypatch.setattr(coordenadas, "PASTA_DADOS", tmp_path)
    assert carregar_metadados() == {}
    assert carregar_metadados(tmp_path / "ausente.xlsx") == {}


def test_metadados_identificados_por_id_e_incluem_pendencias(tmp_path):
    caminho = salvar_referencia(tmp_path, [
        ["008391.0", "Confirmada", "https://fonte.example/8391", "Conversão DMS"],
        [9730, "Pendente", "https://fonte.example/9730", "Coordenada ainda não confirmada"],
        ["8391.5", "Confirmada", "fonte incorreta", "ID não inteiro"],
    ], [" ID Barragem ", " STATUS ", "Fonte", "Observacao"])
    metadados = carregar_metadados(caminho)
    assert set(metadados) == {"8391", "9730"}
    assert metadados["9730"] == {
        "Status": "Pendente", "Fonte": "https://fonte.example/9730",
        "Observação": "Coordenada ainda não confirmada",
    }
    assert metadados["8391"]["Observação"] == "Conversão DMS"


def test_metadados_conflitantes_nao_sao_atribuidos_a_barragem(tmp_path):
    caminho = salvar_referencia(tmp_path, [
        [8391, "Confirmada", "fonte A", ""],
        [8391, "Pendente", "fonte B", ""],
        [8391, "Confirmada", "fonte A", ""],
        [8412, "Confirmada", "fonte A", ""],
    ])
    assert set(carregar_metadados(caminho)) == {"8412"}


def test_referencia_fornecida_tem_41_evidencias_e_34_posicoes_confirmadas():
    metadados = carregar_metadados()
    coordenadas_csv = carregar_coordenadas()
    assert len(metadados) == 41
    assert len(coordenadas_csv) == 34
    pendentes = {
        identificador for identificador, registro in metadados.items()
        if "pendente" in registro["Status"].lower()
    }
    assert len(pendentes) == 7
    assert not pendentes.intersection(coordenadas_csv)
    assert all(registro["Fonte"] and registro["Observação"] for registro in metadados.values())

    # A referência contém as mesmas posições: o enriquecimento continua usando
    # o CSV e não substitui posições válidas pela fonte de metadados.
    caminho = coordenadas.PASTA_DADOS / "referencias" / "coordenadas_barragens_preenchidas.xlsx"
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = livro["Coordenadas"]
        aba.reset_dimensions()
        linhas = aba.iter_rows(values_only=True)
        next(linhas)
        pares = {
            coordenadas.normalizar_id(linha[0]): {"Latitude": linha[2], "Longitude": linha[3]}
            for linha in linhas if linha[2] is not None and linha[3] is not None
        }
        assert pares == coordenadas_csv
    finally:
        livro.close()


def test_enriquecimento_padrao_preserva_fontes_mesmo_sem_coordenadas():
    originais = [{"IDBarragem": 8391}, {"IDBarragem": 9730}]
    enriquecidos = enriquecer_registros(originais)
    assert enriquecidos[0]["Latitude"] == carregar_coordenadas()["8391"]["Latitude"]
    assert "pendente" in enriquecidos[1]["Status"].lower()
    assert "Latitude" not in enriquecidos[1]
    assert enriquecidos[1]["Fonte"]
    assert originais == [{"IDBarragem": 8391}, {"IDBarragem": 9730}]


def test_coordenadas_explicitas_nao_recebem_metadados_de_outra_fonte():
    registro = {"IDBarragem": 8391, "Status": "Fonte própria"}
    enriquecido = enriquecer_registros(
        [registro], {"8391": {"Latitude": -10, "Longitude": -40}}
    )[0]
    assert enriquecido == {"IDBarragem": 8391, "Status": "Fonte própria", "Latitude": -10.0, "Longitude": -40.0}
    assert registro == {"IDBarragem": 8391, "Status": "Fonte própria"}
