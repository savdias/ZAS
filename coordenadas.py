"""Coordenadas opcionais das barragens, sem geocodificação ou acesso à rede."""

import csv
import html
import json
import math
from collections.abc import Iterable, Mapping
from decimal import Decimal, InvalidOperation
from pathlib import Path
import unicodedata

import plotly.graph_objects as go
from openpyxl import load_workbook


PASTA_DADOS = Path(__file__).resolve().parent / "dados"


def normalizar_id(valor):
    """Converte IDs inteiros da planilha/CSV para a mesma chave textual."""
    if valor is None or isinstance(valor, bool):
        return None
    try:
        numero = Decimal(str(valor).strip())
    except (InvalidOperation, ValueError):
        return None
    if not numero.is_finite() or numero != numero.to_integral_value():
        return None
    return str(int(numero))


def _coordenada(valor, limite):
    if valor is None or isinstance(valor, bool):
        return None
    try:
        numero = float(str(valor).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(numero) or not -limite <= numero <= limite:
        return None
    return numero


def _par_coordenadas(registro):
    latitude = _coordenada(registro.get("Latitude"), 90)
    longitude = _coordenada(registro.get("Longitude"), 180)
    if latitude is None or longitude is None:
        return None
    return {"Latitude": latitude, "Longitude": longitude}


def carregar_coordenadas(caminho=PASTA_DADOS / "coordenadas_barragens.csv"):
    """Lê pares válidos do CSV UTF-8 separado por ponto e vírgula.

    Um arquivo ausente ou somente com coordenadas vazias resulta em ``{}``.
    Nomes de barragens não são usados para associar localizações. IDs repetidos
    com coordenadas conflitantes também são ignorados para evitar uma associação
    incorreta.
    """
    coordenadas = {}
    ids_ambiguos = set()
    try:
        arquivo = Path(caminho).open(encoding="utf-8-sig", newline="")
    except FileNotFoundError:
        return coordenadas

    with arquivo:
        leitor = csv.DictReader(arquivo, delimiter=";")
        if leitor.fieldnames:
            leitor.fieldnames = [campo.strip() for campo in leitor.fieldnames]
        for linha in leitor:
            identificador = normalizar_id(linha.get("IDBarragem"))
            par = _par_coordenadas(linha)
            if identificador is None or par is None or identificador in ids_ambiguos:
                continue
            if identificador in coordenadas and coordenadas[identificador] != par:
                del coordenadas[identificador]
                ids_ambiguos.add(identificador)
            else:
                coordenadas[identificador] = par
    return coordenadas


def _normalizar_cabecalho(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    return "".join(letra for letra in texto if letra.isalnum()).lower()


def _identificador_registro(registro):
    if "IDBarragem" in registro:
        return normalizar_id(registro["IDBarragem"])
    for nome, valor in registro.items():
        if _normalizar_cabecalho(nome) in {"idbarragem", "codigodabarragem", "codigobarragem"}:
            return normalizar_id(valor)
    return None


def carregar_metadados(caminho=None):
    """Lê status, fonte e observação do arquivo de referência, por ID.

    A aba ``Coordenadas`` é a única usada. Os metadados também são preservados
    para registros sem posição confirmada. Esta leitura é local e não usa as
    URLs das fontes para consultar ou geocodificar estruturas.
    """
    if caminho is None:
        caminho = PASTA_DADOS / "referencias" / "coordenadas_barragens_preenchidas.xlsx"
    try:
        planilha = load_workbook(caminho, read_only=True, data_only=True)
    except FileNotFoundError:
        return {}

    metadados = {}
    ids_ambiguos = set()
    try:
        if "Coordenadas" not in planilha.sheetnames:
            return metadados
        aba = planilha["Coordenadas"]
        # Algumas planilhas fornecidas não contêm dimensões no XML.
        aba.reset_dimensions()
        linhas = aba.iter_rows(values_only=True)
        cabecalho = next(linhas, ())
        indices = {
            _normalizar_cabecalho(nome): indice
            for indice, nome in enumerate(cabecalho)
            if nome is not None
        }
        indice_id = indices.get("idbarragem")
        if indice_id is None:
            return metadados
        for linha in linhas:
            if indice_id >= len(linha):
                continue
            identificador = normalizar_id(linha[indice_id])
            if identificador is None or identificador in ids_ambiguos:
                continue
            campos = {}
            for nome in ("Status", "Fonte", "Observação"):
                indice = indices.get(_normalizar_cabecalho(nome))
                valor = linha[indice] if indice is not None and indice < len(linha) else None
                campos[nome] = str(valor).strip() if valor is not None else ""
            if identificador in metadados and metadados[identificador] != campos:
                del metadados[identificador]
                ids_ambiguos.add(identificador)
            else:
                metadados[identificador] = campos
    finally:
        planilha.close()
    return metadados


def enriquecer_registros(registros: Iterable[Mapping], coordenadas=None):
    """Cria cópias e acrescenta coordenadas e suas evidências pelo IDBarragem.

    O carregamento padrão usa o CSV para posições e o XLSX para metadados.
    Coordenadas fornecidas explicitamente mantêm o comportamento anterior,
    sem acrescentar metadados de uma outra fonte.
    """
    metadados = {}
    if coordenadas is None:
        coordenadas = carregar_coordenadas()
        metadados = carregar_metadados()
    enriquecidos = []
    for registro in registros:
        copia = dict(registro)
        identificador = _identificador_registro(registro)
        if identificador in metadados:
            copia.update(metadados[identificador])
        coordenada = coordenadas.get(identificador) if identificador is not None else None
        if coordenada is not None:
            par = _par_coordenadas(coordenada)
            if par is not None:
                copia.update(par)
        enriquecidos.append(copia)
    return enriquecidos


def _adicionar_contorno_brasil(figura):
    """Usa somente o contorno GeoJSON local; sua ausência não bloqueia o mapa."""
    try:
        with (PASTA_DADOS / "brasil.geojson").open(encoding="utf-8") as arquivo:
            feature = json.load(arquivo)
    except (OSError, json.JSONDecodeError):
        return

    geometria = feature.get("geometry") or {}
    tipo = geometria.get("type")
    poligonos = geometria.get("coordinates") or []
    if tipo == "Polygon":
        poligonos = [poligonos]
    elif tipo != "MultiPolygon":
        return

    for poligono in poligonos:
        if not poligono:
            continue
        contorno = poligono[0]
        if len(contorno) < 3:
            continue
        pontos = []
        for ponto in contorno:
            if not isinstance(ponto, (list, tuple)) or len(ponto) < 2:
                break
            longitude = _coordenada(ponto[0], 180)
            latitude = _coordenada(ponto[1], 90)
            if longitude is None or latitude is None:
                break
            pontos.append((longitude, latitude))
        else:
            figura.add_trace(go.Scatter(
                x=[ponto[0] for ponto in pontos],
                y=[ponto[1] for ponto in pontos],
                mode="lines",
                name="Brasil",
                line={"color": "#9EB7C4", "width": 1},
                fill="toself",
                fillcolor="#EAF2F4",
                hoverinfo="skip",
            ))


def mapa_barragens(
    registros: Iterable[Mapping], exibir_barragens=True, exibir_contorno=True
):
    """Mostra posições geográficas sem baixar mapas, tiles ou topojson.

    Longitude e latitude são plotadas em eixos cartesianos com proporção
    aproximada para o Brasil e contorno Natural Earth armazenado localmente.
    Sem um par de coordenadas válido, nenhum ponto de barragem é criado.
    """
    localizados = []
    total = 0
    for registro in registros:
        total += 1
        par = _par_coordenadas(registro)
        if par is not None:
            localizados.append((registro, par))

    figura = go.Figure()
    if exibir_contorno:
        _adicionar_contorno_brasil(figura)
    if localizados and exibir_barragens:
        dados_hover = [
            [
                html.escape(str(registro.get("NMBarragem") or "Barragem")),
                html.escape(str(
                    registro.get("NMMineradora")
                    or registro.get("NMEmpreendedor")
                    or "Não informado"
                )),
                _identificador_registro(registro) or "",
            ]
            for registro, _ in localizados
        ]
        figura.add_trace(go.Scatter(
            x=[par["Longitude"] for _, par in localizados],
            y=[par["Latitude"] for _, par in localizados],
            mode="markers",
            name="Barragens",
            marker={"size": 11, "color": "#0F7637", "line": {"width": 1, "color": "white"}},
            customdata=dados_hover,
            hovertemplate=(
                "<b>%{customdata[0]}</b><br>"
                "Mineradora: %{customdata[1]}<br>ID: %{customdata[2]}<br>"
                "Latitude: %{y:.5f}<br>Longitude: %{x:.5f}<extra></extra>"
            ),
        ))
    elif not exibir_barragens:
        figura.add_annotation(
            text="Camada de barragens desativada.",
            x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False,
            font={"size": 15, "color": "#64748B"},
            bgcolor="white",
        )
    else:
        figura.add_annotation(
            text="Não há coordenadas válidas para as barragens selecionadas.",
            x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False,
            font={"size": 15, "color": "#64748B"},
            bgcolor="white",
        )

    figura.add_annotation(
        text=f"{len(localizados)} de {total} registros com coordenadas",
        x=1, y=1.04, xref="paper", yref="paper", xanchor="right",
        showarrow=False, font={"size": 12, "color": "#64748B"},
    )

    longitudes = [par["Longitude"] for _, par in localizados]
    latitudes = [par["Latitude"] for _, par in localizados]
    figura.update_layout(
        template="plotly_white",
        title={"text": "Localização das barragens", "x": 0.02, "font": {"size": 18}},
        height=480,
        margin={"l": 55, "r": 25, "t": 80, "b": 55},
        showlegend=False,
        clickmode="event+select",
        dragmode="zoom",
        uirevision="coordenadas-barragens",
        xaxis={
            "title": "Longitude (°)",
            "range": [min([-75] + [v - 1 for v in longitudes]), max([-33] + [v + 1 for v in longitudes])],
            "zeroline": False,
        },
        yaxis={
            "title": "Latitude (°)",
            "range": [min([-35] + [v - 1 for v in latitudes]), max([7] + [v + 1 for v in latitudes])],
            "scaleanchor": "x",
            "scaleratio": 1 / math.cos(math.radians(-14)),
            "zeroline": False,
        },
    )
    return figura
