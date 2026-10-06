"""Leitura das planilhas de barragens e cálculo dos indicadores do painel.

Uma planilha local é carregada quando existe em ``dados/``. A ausência da
planilha não impede a inicialização do app. Uploads podem ser lidos com
``carregar_planilha(BytesIO(...))`` sem alterar os dados globais.
"""

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import math
import re
import unicodedata
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


PASTA_PROJETO = Path(__file__).resolve().parent
PASTA_DADOS = PASTA_PROJETO / "dados"

ALIASES_COLUNAS = {
    "id": ["IDBarragem", "ID Barragem", "Código da Barragem", "Codigo Barragem"],
    "barragens": ["NMBarragem", "Nome Barragem", "Nome da Barragem", "Barragem"],
    "estados": ["UF", "SGUF", "Sigla UF", "Estado", "Unidade Federativa"],
    "municipios": [
        "NMMunicipio", "NM Municipio", "Município", "Nome Municipio",
        "Nome do Município", "Município da Barragem",
    ],
    "empreendedores": [
        "NMMineradora", "Nome Mineradora", "Mineradora", "Empreendedor",
        "Empresa", "Razão Social",
    ],
    "trabalhadores": [
        "TrabalhadoresZAS", "Trabalhadores ZAS", "Trabalhadores na ZAS",
        "Quantidade de Trabalhadores na ZAS", "Trabalhadores",
    ],
    "comunidade": [
        "ExisteComunidadeZAS", "Existe Comunidade ZAS", "Comunidade ZAS",
        "Comunidade na ZAS",
    ],
}


def normalizar(texto):
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(texto).strip())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", re.sub(r"[_\-]+", " ", texto.lower())).strip()


def converter_numero(valor):
    """Converte números finitos; textos seguem a notação brasileira.

    Aceita também decimais com ponto e a notação mista ``1,234.56``.
    Um ponto com grupos de três dígitos, como ``1.234``, é de milhar.
    Respostas categóricas, booleanos, NaN e infinito não são contagens.
    """
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        try:
            numero = float(valor)
        except (OverflowError, ValueError):
            return None
        return numero if math.isfinite(numero) else None
    texto = re.sub(r"\s+", "", str(valor))
    if not texto:
        return None
    if texto.startswith("(") and texto.endswith(")"):
        texto = "-" + texto[1:-1]
    if "," in texto and "." in texto:
        decimal = "," if texto.rfind(",") > texto.rfind(".") else "."
        milhar = "." if decimal == "," else ","
        if texto.count(decimal) != 1:
            return None
        inteiro, fracao = texto.rsplit(decimal, 1)
        if not re.fullmatch(r"[+-]?\d{1,3}(?:" + re.escape(milhar) + r"\d{3})+", inteiro):
            return None
        texto = inteiro.replace(milhar, "") + "." + fracao
    elif "," in texto:
        if texto.count(",") != 1:
            return None
        texto = texto.replace(",", ".")
    elif re.fullmatch(r"[+-]?[1-9]\d{0,2}(?:\.\d{3})+", texto):
        texto = texto.replace(".", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", texto):
        return None
    try:
        numero = float(texto)
    except (OverflowError, ValueError):
        return None
    return numero if math.isfinite(numero) else None


def formatar_numero(numero):
    numero = converter_numero(numero)
    if numero is None:
        return "—"
    if numero.is_integer():
        return f"{int(numero):,}".replace(",", ".")
    texto = f"{numero:,.2f}".rstrip("0").rstrip(".")
    return texto.translate(str.maketrans({",": ".", ".": ","}))


def localizar_planilha():
    arquivos = sorted(
        (p for p in PASTA_DADOS.glob("*")
         if p.is_file() and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")
         and "coordenad" not in normalizar(p.stem)),
        key=lambda p: p.name.casefold(),
    )
    if not arquivos:
        raise FileNotFoundError(f"Nenhuma planilha XLSX encontrada em {PASTA_DADOS}")
    prioritarios = [p for p in arquivos if any(
        termo in normalizar(p.stem) for termo in ("barragem", "anm", "air")
    )]
    return (prioritarios or arquivos)[0]


def localizar_coluna(cabecalhos, possibilidades):
    mapa = {}
    for cabecalho in cabecalhos:
        if normalizar(cabecalho):
            mapa.setdefault(normalizar(cabecalho), cabecalho)
    for possibilidade in possibilidades:
        if normalizar(possibilidade) in mapa:
            return mapa[normalizar(possibilidade)]
    for possibilidade in possibilidades:
        chave = normalizar(possibilidade)
        if len(chave) < 3:
            continue
        # Limites de palavra evitam "Barragem" dentro de "IDBarragem".
        correspondencias = [
            (chave_planilha, original) for chave_planilha, original in mapa.items()
            if re.search(r"(?<!\w)" + re.escape(chave) + r"(?!\w)", chave_planilha)
        ]
        if correspondencias:
            return min(correspondencias, key=lambda item: len(item[0]))[1]
    return None


def identificar_colunas(cabecalhos):
    donos_exatos = {normalizar(alias): chave
                    for chave, aliases in ALIASES_COLUNAS.items() for alias in aliases}
    return {
        chave: localizar_coluna(
            [c for c in cabecalhos if donos_exatos.get(normalizar(c), chave) == chave],
            aliases,
        )
        for chave, aliases in ALIASES_COLUNAS.items()
    }


def _avaliar_cabecalho(ws):
    melhor = (None, (0, 0, 0), None)
    limite = min(ws.max_row or 30, 30)
    for numero, valores in enumerate(ws.iter_rows(max_row=limite, values_only=True), 1):
        textos = [str(v).strip() if isinstance(v, str) else "" for v in valores]
        colunas = identificar_colunas(textos)
        if not (colunas["id"] or colunas["barragens"]):
            continue
        exatas = sum(
            normalizar(colunas[chave]) in {normalizar(a) for a in aliases}
            for chave, aliases in ALIASES_COLUNAS.items() if colunas[chave]
        )
        reconhecidas = sum(v is not None for v in colunas.values())
        pontos = (exatas * 20 + (reconhecidas - exatas) * 5,
                  reconhecidas, sum(bool(t) for t in textos))
        if pontos[0] and pontos > melhor[1]:
            melhor = (numero, pontos, valores)
    return melhor


def descobrir_cabecalho(ws):
    linha, _, _ = _avaliar_cabecalho(ws)
    if linha is None:
        raise ValueError(f"Nenhum cabeçalho de barragens reconhecido na aba {ws.title!r}.")
    return linha


def _cabecalhos_unicos(valores):
    cabecalhos = []
    usados = set()
    for valor in valores:
        original = "" if valor is None else str(valor).strip()
        nome = original
        numero = 2
        while nome and normalizar(nome) in usados:
            nome = f"{original} ({numero})"
            numero += 1
        if nome:
            usados.add(normalizar(nome))
        cabecalhos.append(nome)
    return cabecalhos


def _valor_json(valor):
    if isinstance(valor, (datetime, date, time)):
        return valor.isoformat()
    if isinstance(valor, str):
        return valor.strip() or None
    if isinstance(valor, float) and not math.isfinite(valor):
        return None
    return valor


def carregar_planilha(arquivo=None):
    """Retorna ``(cabecalhos, registros)`` de um caminho ou arquivo binário.

    Seleciona a aba pelos campos do domínio, preserva cabeçalhos duplicados e
    fecha o workbook mesmo quando houver erro. Não altera os dados globais.
    """
    if arquivo is None:
        arquivo = localizar_planilha()
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    try:
        candidatos = [(ws, *_avaliar_cabecalho(ws)) for ws in workbook.worksheets]
        candidatos = [item for item in candidatos if item[1] is not None]
        if not candidatos:
            raise ValueError(
                "A planilha não possui cabeçalhos reconhecidos de barragens: "
                "é necessário informar o ID ou o nome da barragem."
            )
        ws, linha_cabecalho, _, valores = max(candidatos, key=lambda item: item[2])
        cabecalhos = _cabecalhos_unicos(valores)
        registros = []
        for linha in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
            registro = {cabecalho: _valor_json(valor)
                        for cabecalho, valor in zip(cabecalhos, linha) if cabecalho}
            if any(v is not None and v != "" for v in registro.values()):
                registros.append(registro)
        return cabecalhos, registros
    finally:
        workbook.close()


def valores_unicos(coluna, registros=None):
    if not coluna:
        return set()
    if registros is None:
        registros = REGISTROS
    return {str(r[coluna]).strip() for r in registros
            if r.get(coluna) is not None and str(r[coluna]).strip()}


def somar_coluna(coluna, registros=None):
    if not coluna:
        return None
    if registros is None:
        registros = REGISTROS
    numeros = [converter_numero(r.get(coluna)) for r in registros]
    numeros = [numero for numero in numeros if numero is not None]
    return math.fsum(numeros) if numeros else None


def contar_respostas_sim(registros, coluna):
    """Conta barragens com resposta afirmativa, sem duplicar sua identidade.

    Números são quantidades, não respostas Sim/Não. Sem ID ou nome disponível,
    cada registro representa uma resposta. Ausência de resposta válida é None.
    """
    if not coluna:
        return None
    sim = {"sim", "s", "yes", "true"}
    nao = {"nao", "n", "no", "false"}
    cabecalhos = list(dict.fromkeys(c for r in registros for c in r))
    colunas = identificar_colunas(cabecalhos)
    afirmativas = set()
    encontrou = False
    for indice, registro in enumerate(registros):
        resposta = normalizar(registro.get(coluna))
        if resposta not in sim | nao:
            continue
        encontrou = True
        if resposta in sim:
            afirmativas.add(_identidade_barragem(registro, colunas) or ("linha", indice))
    return len(afirmativas) if encontrou else None


def _chave_entidade(valor):
    if isinstance(valor, float) and math.isfinite(valor) and valor.is_integer():
        valor = int(valor)
    return normalizar(valor)


def _contar_entidades(registros, coluna, contexto=()):
    if not coluna:
        return None
    return len({tuple(_chave_entidade(r.get(c)) for c in contexto if c)
                + (_chave_entidade(r.get(coluna)),)
                for r in registros if _chave_entidade(r.get(coluna))})


def _identidade_barragem(registro, colunas):
    identificador = colunas.get("id")
    if identificador and _chave_entidade(registro.get(identificador)):
        return "id", _chave_entidade(registro.get(identificador))
    barragem = colunas.get("barragens")
    if barragem and _chave_entidade(registro.get(barragem)):
        return ("nome",) + tuple(
            _chave_entidade(registro.get(colunas.get(chave)))
            for chave in ("estados", "municipios", "barragens")
        )
    return None


def calcular_indicadores(registros, colunas=None):
    if colunas is None:
        cabecalhos = list(dict.fromkeys(c for r in registros for c in r))
        colunas = identificar_colunas(cabecalhos)
    estado = colunas.get("estados")
    municipio = colunas.get("municipios")
    barragem = colunas.get("barragens")
    identificador = colunas.get("id")
    identidades = {_identidade_barragem(registro, colunas) for registro in registros}
    identidades.discard(None)
    totais = {
        "barragens": len(identidades) if identificador or barragem else None,
        "estados": _contar_entidades(registros, estado),
        "municipios": _contar_entidades(registros, municipio, (estado,)),
        "empreendedores": _contar_entidades(registros, colunas.get("empreendedores")),
        "trabalhadores": somar_coluna(colunas.get("trabalhadores"), registros),
        "registros": len(registros),
    }
    return {chave: formatar_numero(valor) for chave, valor in totais.items()}


CABECALHOS, REGISTROS = [], []
ERRO_DADOS = None
try:
    CABECALHOS, REGISTROS = carregar_planilha()
except (OSError, ValueError, BadZipFile, InvalidFileException, ParseError) as erro:
    ERRO_DADOS = str(erro)

COLUNAS = identificar_colunas(CABECALHOS)
COLUNA_ID = COLUNAS["id"]
COLUNA_BARRAGEM = COLUNAS["barragens"]
COLUNA_ESTADO = COLUNAS["estados"]
COLUNA_MUNICIPIO = COLUNAS["municipios"]
COLUNA_EMPREENDEDOR = COLUNAS["empreendedores"]
COLUNA_TRABALHADORES = COLUNAS["trabalhadores"]
COLUNA_COMUNIDADE = COLUNAS["comunidade"]

BARRAGENS = valores_unicos(COLUNA_BARRAGEM)
ESTADOS = valores_unicos(COLUNA_ESTADO)
MUNICIPIOS = valores_unicos(COLUNA_MUNICIPIO)
EMPREENDEDORES = valores_unicos(COLUNA_EMPREENDEDOR)
INDICADORES = calcular_indicadores(REGISTROS, COLUNAS)
TOTAL_BARRAGENS = converter_numero(INDICADORES["barragens"])
TOTAL_ESTADOS = converter_numero(INDICADORES["estados"])
TOTAL_MUNICIPIOS = converter_numero(INDICADORES["municipios"])
TOTAL_EMPREENDEDORES = converter_numero(INDICADORES["empreendedores"])
TOTAL_TRABALHADORES = somar_coluna(COLUNA_TRABALHADORES)
